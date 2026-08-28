"""Carregamento de deparas a partir de .xlsx (openpyxl) ou .csv."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from agent.domain.classificacao_nf.deparas import TabelasDepara, carregar_tabela_csv


def carregar_tabela_xlsx(caminho: str | Path, *, sheet: int | str = 0) -> list[dict[str, str]]:
    """Primeira linha = cabeçalho. Células viram str.strip()."""
    from openpyxl import load_workbook

    path = Path(caminho)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[int(sheet)]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return []
        chaves = [str(c or "").strip() for c in header]
        out: list[dict[str, str]] = []
        for row in rows:
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            item: dict[str, str] = {}
            for i, chave in enumerate(chaves):
                if not chave:
                    continue
                valor = row[i] if i < len(row) else ""
                item[chave] = "" if valor is None else str(valor).strip()
            out.append(item)
        return out
    finally:
        wb.close()


def carregar_tabela_arquivo(caminho: str | Path) -> list[dict[str, str]]:
    path = Path(caminho)
    sufixo = path.suffix.casefold()
    if sufixo == ".csv":
        return carregar_tabela_csv(path)
    if sufixo in (".xlsx", ".xlsm"):
        return carregar_tabela_xlsx(path)
    raise ValueError(f"formato de depara não suportado: {path.suffix}")


def carregar_tabelas_depara(
    *,
    natureza_despesa: str | Path | None = None,
    codigo_servico: str | Path | None = None,
    natureza_rendimento: str | Path | None = None,
    vencimento_especial: str | Path | None = None,
) -> TabelasDepara:
    """Monta TabelasDepara a partir de paths (vazios = tabela vazia)."""

    def _load(path: str | Path | None) -> list[dict[str, Any]]:
        if not path or not str(path).strip():
            return []
        return carregar_tabela_arquivo(path)

    return TabelasDepara(
        natureza_despesa=_load(natureza_despesa),
        codigo_servico=_load(codigo_servico),
        natureza_rendimento=_load(natureza_rendimento),
        vencimento_especial=_load(vencimento_especial),
    )


def carregar_tabelas_depara_de_config() -> TabelasDepara:
    from agent import config

    return carregar_tabelas_depara(
        natureza_despesa=config.DEPARA_NATUREZA_DESPESA_PATH or None,
        codigo_servico=config.DEPARA_CODIGO_SERVICO_PATH or None,
        natureza_rendimento=config.DEPARA_NATUREZA_RENDIMENTO_PATH or None,
        vencimento_especial=config.DEPARA_VENCIMENTO_ESPECIAL_PATH or None,
    )
