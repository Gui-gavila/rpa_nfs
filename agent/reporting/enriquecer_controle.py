# -*- coding: utf-8 -*-
"""Enriquecimento offline de ControleClassificacao com JSON LYNN (P1/P2).

Sem chamadas a Protheus ou LYNN — apenas Excel ops + ``{stem}.json`` locais.
Schema alvo = colunas de ``ControleClassificacao-Avaliacao.xlsx``.
"""

from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping

from agent.integrations.lynn.offline import LynnOfflinePayload, carregar_lynn_arquivo

logger = logging.getLogger(__name__)

# Ordem canónica = cabeçalho de ControleClassificacao-Avaliacao.xlsx
COLUNAS_AVALIACAO: tuple[str, ...] = (
    "COD_FILIAL",
    "NOME_FILIAL",
    "CNPJ_FILIAL",
    "COD_FORNECEDOR",
    "NOME_FORNECEDOR",
    "CNPJ_FORNECEDOR",
    "COD_LOJA",
    "NIVEL_PJ",
    "NUMERO_NF",
    "DATA_EMISSAO_NF",
    "VALOR_TOTAL_NF",
    "SERIE_NF",
    "COD_ENTIDADE",
    "NUMERO_NF_PDF",
    "TIPO_NF_PDF",
    "DATA_EMISSAO_NF_PDF",
    "VALOR_TOTAL_NF_PDF",
    "CNPJ_TOMADOR_PDF",
    "CNPJ_PRESTADOR_PDF",
    "COD_TRIBUTACAO",
    "ISSQN",
    "IRRF",
    "PIS",
    "COFINS",
    "CSLL",
    "LAYOUT",
    "MEI",
    "RETENCAO",
    "ISSQN_RETIDO",
    "NATUREZA_DESPESA",
    "NATUREZA_RENDIMENTO",
    "COD_RETENCAO_IRRF",
    "COD_RETENCAO_PCC",
    "STATUS",
    "MOTIVO",
    "PDF_RECUPERADO",
    "LYNN_PROCESSADO",
    "STATUS_LYNN",
    "STATUS_CLASSIFICACAO_PROTHEUS",
    "TIMESTAMP ENTRADA",
    "TIMESTAMP SAIDA",
)


def _s(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    text = str(v).strip()
    if text in {"-", "None", "null"}:
        return ""
    return text


def _dig(obj: Mapping[str, Any] | None, *path: str) -> Any:
    cur: Any = obj
    for key in path:
        if not isinstance(cur, Mapping):
            return None
        cur = cur.get(key)
    return cur


def _primeiro(*vals: Any) -> str:
    for v in vals:
        s = _s(v)
        if s:
            return s
    return ""


def _boolish(v: Any) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    s = str(v).strip().casefold()
    if s in {"true", "1", "sim", "yes", "s"}:
        return "Sim"
    if s in {"false", "0", "nao", "não", "no", "n"}:
        return "Não"
    return _s(v)


def mapear_linha_avaliacao(
    *,
    ops: Mapping[str, Any] | None = None,
    extracao: Mapping[str, Any] | None = None,
    classificacao: Mapping[str, Any] | None = None,
    timestamp_entrada: str = "",
    timestamp_saida: str = "",
) -> dict[str, str]:
    """De-para: linha ops (COLUNAS_RELATORIO_NF) + blocos Lynn → schema Avaliação.

    Regras:
    - Campos ERP/ops: ``FILIAL``→``COD_FILIAL``, ``SERIE_PDF``→``SERIE_NF``, etc.
    - Campos PDF/Lynn: preferir ``classificacao`` / ``extracao`` quando presentes.
    - ``STATUS`` / ``MOTIVO``: prevalece ops (Carol).
    - ``STATUS_LYNN``: status textual do agente Lynn.
    - ``STATUS_CLASSIFICACAO_PROTHEUS``: preenchido após UI/Protheus OK.
    - ``PDF_RECUPERADO`` / ``LYNN_PROCESSADO``: flags Sim/Não de gate.
    - ``TIMESTAMP ENTRADA`` / ``TIMESTAMP SAIDA``: horários de negócio
      (PDF recuperado / classificada no Protheus); não derivar de mtime.
    """
    ops = dict(ops or {})
    ex = extracao if isinstance(extracao, Mapping) else None
    clf = classificacao if isinstance(classificacao, Mapping) else None

    nfse = _dig(ex, "nfse_info") if ex else None
    prest = _dig(ex, "x_prestador") if ex else None
    tom = _dig(ex, "x_tomador") if ex else None
    imp = _dig(ex, "x_impostos_valores") if ex else None

    cnpj_tomador = _primeiro(_dig(tom, "x_cnpj_tomador") if tom else None)
    cnpj_prestador = _primeiro(_dig(prest, "x_cnpj_prestador") if prest else None)
    nome_tomador = _primeiro(_dig(tom, "x_nome_tomador") if tom else None)

    status_lynn = _primeiro(
        ops.get("STATUS_LYNN"),
        _dig(clf, "STATUS") if clf else None,
        ops.get("STATUS_PROCESSAMENTO"),  # legado
    )
    mei = _boolish(_dig(clf, "MEI") if clf else None) or _boolish(
        _dig(prest, "x_mei") if prest else None
    )
    iss_ret = _boolish(_dig(clf, "ISSQN_RETIDO") if clf else None) or _boolish(
        _dig(tom, "x_issqn_retido_tomador") if tom else None
    )
    retencao = _boolish(_dig(nfse, "x_tipo_retencao") if nfse else None)
    if not retencao and iss_ret == "Sim":
        retencao = "Sim"

    tipo_nf_pdf = _boolish(_dig(clf, "TIPO_NF_PDF") if clf else None)
    if not tipo_nf_pdf:
        tipo_nf_pdf = _boolish(_dig(nfse, "x_nfse") if nfse else None)

    lynn_proc = _primeiro(ops.get("LYNN_PROCESSADO"))
    if not lynn_proc and status_lynn:
        lynn_proc = "Sim"
    pdf_rec = _primeiro(ops.get("PDF_RECUPERADO"))
    if not pdf_rec and _primeiro(ops.get("COD_ENTIDADE"), ops.get("COD_OBJETO")):
        pdf_rec = "Sim"

    linha = {
        "COD_FILIAL": _primeiro(ops.get("FILIAL"), ops.get("COD_FILIAL")),
        "NOME_FILIAL": _primeiro(ops.get("NOME_FILIAL"), nome_tomador),
        "CNPJ_FILIAL": _primeiro(ops.get("CNPJ_FILIAL"), cnpj_tomador),
        "COD_FORNECEDOR": _primeiro(ops.get("COD_FORNECEDOR")),
        "NOME_FORNECEDOR": _primeiro(
            ops.get("NOME_FORNECEDOR"),
            _dig(prest, "x_nome_prestador") if prest else None,
        ),
        "CNPJ_FORNECEDOR": _primeiro(ops.get("CNPJ_FORNECEDOR"), cnpj_prestador),
        "COD_LOJA": _primeiro(ops.get("COD_LOJA")),
        "NIVEL_PJ": _primeiro(ops.get("NIVEL_PJ")),
        "NUMERO_NF": _primeiro(ops.get("NUMERO_NF")),
        "DATA_EMISSAO_NF": _primeiro(ops.get("DATA_EMISSAO_NF")),
        "VALOR_TOTAL_NF": _primeiro(ops.get("VALOR_TOTAL_NF")),
        "SERIE_NF": _primeiro(
            ops.get("SERIE_PDF"),
            ops.get("SERIE_NF"),
            _dig(nfse, "x_dps_serie") if nfse else None,
        ),
        "COD_ENTIDADE": _primeiro(ops.get("COD_ENTIDADE"), ops.get("COD_OBJETO")),
        "NUMERO_NF_PDF": _primeiro(
            _dig(nfse, "x_numero_nfs") if nfse else None,
            ops.get("NUMERO_NF_PDF"),
            ops.get("NUMERO_NF"),
        ),
        "TIPO_NF_PDF": _primeiro(tipo_nf_pdf, ops.get("TIPO_NF_PDF")),
        "DATA_EMISSAO_NF_PDF": _primeiro(
            _dig(nfse, "x_data_emissao") if nfse else None,
            ops.get("DATA_EMISSAO_NF_PDF"),
        ),
        "VALOR_TOTAL_NF_PDF": _primeiro(
            _dig(nfse, "x_valor_total_nfs") if nfse else None,
            ops.get("VALOR_TOTAL_NF_PDF"),
        ),
        "CNPJ_TOMADOR_PDF": cnpj_tomador,
        "CNPJ_PRESTADOR_PDF": cnpj_prestador,
        "COD_TRIBUTACAO": _primeiro(
            _dig(clf, "COD_TRIBUTACAO") if clf else None,
            _dig(nfse, "x_cod_servico") if nfse else None,
            ops.get("COD_SERVICO"),
            ops.get("COD_TRIBUTACAO"),
        ),
        "ISSQN": _primeiro(
            _dig(imp, "x_valor_issqn") if imp else None,
            ops.get("ISSQN"),
        ),
        "IRRF": _primeiro(
            _dig(imp, "x_valor_irrf") if imp else None,
            ops.get("IRRF"),
        ),
        "PIS": _primeiro(
            _dig(imp, "x_valor_pis") if imp else None,
            ops.get("PIS"),
        ),
        "COFINS": _primeiro(
            _dig(imp, "x_valor_cofins") if imp else None,
            ops.get("COFINS"),
        ),
        "CSLL": _primeiro(
            _dig(imp, "x_valor_csll") if imp else None,
            ops.get("CSLL"),
        ),
        "LAYOUT": _primeiro(_dig(clf, "LAYOUT") if clf else None, ops.get("LAYOUT")),
        "MEI": mei,
        "RETENCAO": retencao,
        "ISSQN_RETIDO": iss_ret,
        "NATUREZA_DESPESA": _primeiro(
            _dig(clf, "NATUREZA_DESPESA") if clf else None,
            ops.get("NATUREZA_DESPESA"),
        ),
        "NATUREZA_RENDIMENTO": _primeiro(
            _dig(clf, "NATUREZA_RENDIMENTO") if clf else None,
            ops.get("NATUREZA_RENDIMENTO"),
        ),
        "COD_RETENCAO_IRRF": _primeiro(
            _dig(clf, "COD_RETENCAO_IRRF") if clf else None,
            ops.get("COD_RETENCAO_IRRF"),
        ),
        "COD_RETENCAO_PCC": _primeiro(
            _dig(clf, "COD_RETENCAO_PCC") if clf else None,
            ops.get("COD_RETENCAO_PCC"),
        ),
        "STATUS": _primeiro(ops.get("STATUS")),
        "MOTIVO": _primeiro(ops.get("MOTIVO")),
        "PDF_RECUPERADO": pdf_rec,
        "LYNN_PROCESSADO": lynn_proc,
        "STATUS_LYNN": status_lynn,
        "STATUS_CLASSIFICACAO_PROTHEUS": _primeiro(
            ops.get("STATUS_CLASSIFICACAO_PROTHEUS")
        ),
        "TIMESTAMP ENTRADA": _s(timestamp_entrada),
        "TIMESTAMP SAIDA": _s(timestamp_saida),
    }
    return {c: _s(linha.get(c, "")) for c in COLUNAS_AVALIACAO}


def stem_ops(ops: Mapping[str, Any]) -> str:
    """``{filial}_{numero}_{fornecedor}`` alinhado ao naming do data-plane."""
    filial = _primeiro(ops.get("FILIAL"), ops.get("COD_FILIAL"))
    numero = _primeiro(ops.get("NUMERO_NF"))
    forn = _primeiro(ops.get("COD_FORNECEDOR"))
    return f"{filial}_{numero}_{forn}"


def ops_de_stem(stem: str) -> dict[str, str]:
    """Reconstrói chaves ops a partir do stem ``filial_numero_fornecedor``."""
    partes = stem.split("_")
    if len(partes) < 3:
        return {"FILIAL": "", "NUMERO_NF": "", "COD_FORNECEDOR": stem}
    filial, numero, forn = partes[0], partes[1], "_".join(partes[2:])
    return {
        "FILIAL": filial,
        "COD_FILIAL": filial,
        "NUMERO_NF": numero,
        "COD_FORNECEDOR": forn,
        "NUMERO_NF_PDF": numero,
    }


def indice_json_lynn(pasta: str | Path) -> dict[str, Path]:
    """Indexa ``*.json`` em ``02_NF_Processadas_Lynn`` (e subpastas) por stem."""
    root = Path(pasta)
    out: dict[str, Path] = {}
    if not root.is_dir():
        return out
    for path in root.rglob("*.json"):
        out[path.stem] = path
    return out


def stems_processadas(pasta: str | Path) -> list[str]:
    """Stems únicos de PDF/JSON na pasta processadas (ordem estável)."""
    root = Path(pasta)
    stems: set[str] = set()
    if not root.is_dir():
        return []
    for path in root.rglob("*"):
        if path.suffix.lower() in {".pdf", ".json"}:
            stems.add(path.stem)
    return sorted(stems)


def _mtime_iso(path: Path) -> str:
    try:
        ts = path.stat().st_mtime
        return datetime.fromtimestamp(ts, tz=timezone.utc).astimezone().isoformat(
            timespec="seconds"
        )
    except OSError:
        return ""


def agora_iso_local() -> str:
    """Data/hora local atual (TIMESTAMP ENTRADA / SAIDA de negócio)."""
    return datetime.now().astimezone().isoformat(timespec="seconds")


def ler_linhas_ops_xlsx(caminho: str | Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    path = Path(caminho)
    if not path.is_file():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    # Já no schema Avaliação? tratar COD_FILIAL como ops.
    linhas: list[dict[str, str]] = []
    for row in rows[1:]:
        if not row or not any(c is not None and str(c).strip() for c in row):
            continue
        item = {
            headers[i]: _s(row[i] if i < len(row) else None)
            for i in range(len(headers))
            if headers[i]
        }
        if "FILIAL" not in item and item.get("COD_FILIAL"):
            item["FILIAL"] = item["COD_FILIAL"]
        if "SERIE_PDF" not in item and item.get("SERIE_NF"):
            item["SERIE_PDF"] = item["SERIE_NF"]
        if "COD_SERVICO" not in item and item.get("COD_TRIBUTACAO"):
            item["COD_SERVICO"] = item["COD_TRIBUTACAO"]
        linhas.append(item)
    return linhas


def gravar_avaliacao_xlsx(caminho: str | Path, linhas: list[dict[str, str]]) -> Path:
    from openpyxl import Workbook

    path = Path(caminho)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "ESTRUTURA BD TRANSACIONAL"
    ws.append(list(COLUNAS_AVALIACAO))
    for linha in linhas:
        ws.append([linha.get(c, "") for c in COLUNAS_AVALIACAO])
    wb.save(path)
    return path


def enriquecer_controle_offline(
    *,
    pasta_processadas: str | Path,
    destino_xlsx: str | Path,
    controle_xlsx: str | Path | None = None,
) -> dict[str, Any]:
    """Gera/atualiza ``ControleClassificacao.xlsx`` no schema Avaliação.

    - Se ``controle_xlsx`` existir: usa essas linhas como base ERP/ops.
    - Se não existir: monta linhas a partir dos stems PDF/JSON em ``pasta_processadas``.
    - Destino padrão operacional: o próprio ``ControleClassificacao.xlsx`` (substitui).
    Sem HTTP Protheus/LYNN.
    """
    pasta = Path(pasta_processadas)
    destino = Path(destino_xlsx)
    fonte = Path(controle_xlsx) if controle_xlsx else None

    ops_linhas: list[dict[str, str]] = []
    origem = "stems"
    if fonte is not None and fonte.is_file():
        # Se destino == fonte, carregar tudo em memória antes de sobrescrever.
        ops_linhas = ler_linhas_ops_xlsx(fonte)
        origem = "excel"
    if not ops_linhas:
        ops_linhas = [ops_de_stem(s) for s in stems_processadas(pasta)]
        origem = "stems"

    indice = indice_json_lynn(pasta)
    enriquecidas = 0
    sem_json = 0
    erros = 0
    saida: list[dict[str, str]] = []

    for ops in ops_linhas:
        stem = stem_ops(ops)
        jpath = indice.get(stem)
        payload: LynnOfflinePayload | None = None
        if jpath is not None:
            try:
                payload = carregar_lynn_arquivo(jpath)
                enriquecidas += 1
            except Exception as e:
                erros += 1
                logger.warning("[enriquecer] falha %s: %s", jpath, e)
                payload = None
        else:
            sem_json += 1

        # Offline não inventa timestamps de negócio (ENTRADA=PDF / SAIDA=Protheus).
        ts_in = _s(ops.get("TIMESTAMP ENTRADA"))
        ts_out = _s(ops.get("TIMESTAMP SAIDA"))

        saida.append(
            mapear_linha_avaliacao(
                ops=ops,
                extracao=payload.extracao if payload else None,
                classificacao=payload.classificacao if payload else None,
                timestamp_entrada=ts_in,
                timestamp_saida=ts_out,
            )
        )

    gravado = gravar_avaliacao_xlsx(destino, saida)
    resumo = {
        "ok": True,
        "destino": str(gravado),
        "origem_linhas": origem,
        "linhas": len(saida),
        "enriquecidas_lynn": enriquecidas,
        "sem_json": sem_json,
        "erros_json": erros,
        "json_indexados": len(indice),
    }
    logger.info("[enriquecer] %s", resumo)
    return resumo
